import re,sys,zipfile,pymupdf
from pptx import Presentation
O=sys.argv[1]
prs=Presentation(f'{O}/01_Today_A_Typical_Day_Student_FINAL.pptx')
slides=[' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).replace('’',"'") for s in prs.slides]
notes=[s.notes_slide.notes_text_frame.text.replace('’',"'") for s in prs.slides]
stu=[re.sub(r'[ \t]*\n[ \t]*',' ',p.get_text()).replace('’',"'") for p in pymupdf.open(f'{O}/02_Today_A_Typical_Day_Student_FINAL.pdf')]
tea=' '.join(p.get_text() for p in pymupdf.open(f'{O}/03_Today_A_Typical_Day_Teacher_Lesson_Plan_FINAL.pdf')).replace('’',"'").replace('\n',' ')
student_all='\n'.join(slides+stu); R=[]
def chk(name,ok,det=''): R.append((name,ok,det))
# character system
chk('No "Lucas" in any file',not re.search('Lucas','\n'.join(slides+notes+stu)+tea))
chk('Deck slides 4–5 show Dr. Sarah',all('Dr. Sarah' in slides[i] for i in (3,4)))
chk('Dr. Ben labelled as Dr. Sarah\'s colleague (deck + PDF)',"DR. SARAH'S COLLEAGUE" in slides[12].upper() and "SARAH'S COLLEAGUE" in stu[3].upper())
chk('Dr. Ben only in Stage 8 slides (13–14)',set(i+1 for i,t in enumerate(slides) if 'BEN' in t.upper())<={13,14},str([i+1 for i,t in enumerate(slides) if 'BEN' in t.upper()]))
# leakage of teacher-only facts into student materials
leaks=['15 patients','She watches TV',"She doesn't work at night","She doesn't work on Sundays",'appointments in the afternoon.','11:00.','bed at 11','8:00 in the evening','20 patients','6:00 in the morning',"doesn't work on Mondays",'Ben works at a hospital']
found=[(l,[k for k,t in enumerate(slides+stu) if l in t]) for l in leaks]
bad=[f for f in found if f[1]]  # strict: no exceptions
chk('Teacher-only facts absent from student slides/PDF',not bad,str(found))
# stage 6 facts in teacher plan and notes
facts=["She sees 15 patients a day","She doesn't work on Sundays","She doesn't work at night","She has appointments in the afternoon","She watches TV after work","She goes to bed at 11:00"]
chk('Sarah facts in teacher plan',all(f in tea for f in facts),[f for f in facts if f not in tea])
nf=["sees 15 patients a day","doesn't work on Sundays","doesn't work at night","has appointments in the afternoon","watches TV after work","goes to bed at 11:00"]
chk('Sarah facts in slide 10 notes',all(f in notes[9] for f in nf))
ben=["works at a hospital","starts at 8:00 in the evening","works at night","sees about 20 patients a night","works on weekends","doesn't work on Mondays","finishes at 6:00 in the morning"]
chk('Ben facts in plan + slide 13 notes',all(b in tea for b in ben) and all(b in notes[12] for b in ben),[b for b in ben if b not in tea or b not in notes[12]])
# timeline consistency
tl=['gets up','arrives at the clinic','sees patients','examines a patient','has lunch','checks medical records','finishes work']
chk('Stage 3 timeline identical: deck slide 5, student PDF p2, plan',all(x in slides[4] for x in tl) and all(x in stu[1] for x in tl) and all(x in tea for x in tl))
chk('No old Stage 6 facts anywhere',not re.search(r'starts at 7:30|lunch at 12:30|finishes at 5:00|works in a hospital','\n'.join(slides+notes+stu)+tea))
# answer keys
chk('Right/Wrong key: plan says 1 wrong 2 wrong 3 right 4 wrong 5 right','1 wrong' in tea and '2 wrong' in tea and '3 right' in tea and '4 wrong' in tea and '5 right' in tea)
chk('Stage 5 wrong instruction removed','looks" correct' not in tea and 'looks” correct' not in tea and 'point back to DOES + HE + WORK and elicit the correction' in tea)
chk('Quick Review key on PDF p6',"1 WORKS" in stu[5] and "2 DOES" in stu[5] and "3 DON'T WORK" in stu[5] and "4 FINISH" in stu[5] and "5 DO" in stu[5])
# pdf blanks
chk('PDF p5 blanks explicit','I see about ______ patients a day.' in stu[4] and "I don't ______________ on weekends." in stu[4] and 'He / She ______________ every day.' in stu[4] and "He / She doesn't ______________." in stu[4], stu[4][:0])
chk('PDF p6 items 2 and 5 have blanks','________ he see patients' in stu[5] and 'What time ________ you start work?' in stu[5])
chk('PDF p4 notes labelled','MY NOTES DURING THE LESSON' in stu[3].upper())
# pronunciation / round2 / delayed correction / exit
chk('Slide 12 USE = "Tell me about a colleague."','Tell me about a colleague.' in slides[11] and 'friend' not in slides[11])
chk('Round 2 note: choose 3–4 questions','Choose 3–4 questions' in notes[13] and 'Choose 3–4 questions' in tea.replace('choose','Choose') )
chk('Slide 16 note: BEFORE SHOWING + one error at a time','BEFORE SHOWING' in notes[15] and 'one error at a time' in notes[15] and 'more than one target error' in notes[15])
chk('Exit vs Round 3 distinguished','2–3 minutes' in notes[14] and '4–5 sentences' in notes[16])
chk('Teaser: "Do not teach this now." in notes + plan','Do not teach this now' in notes[16] and 'do not teach it now' in tea)
# timings
exp=[4,5,7,7,5,7,3,10,5,2]
chk('Stage timings in notes',all(re.search(rf"STAGE {i+1} · [^\n]*{m} min",'\n'.join(notes)) for i,m in enumerate(exp)))
chk('Core 55 min stated in plan','55 minutes' in tea)
chk('17 slides',len(slides)==17)

# FINAL correction checks
chk('Slide 9 sentence 3 = "She doesn\'t work on Fridays."',"3 She doesn't work on Fridays." in slides[8].replace('\n',' ') or "She doesn't work on Fridays." in slides[8])
chk('Old sentence 3 absent from all files',"She doesn't work at night." not in ' '.join(slides+stu) and "at night. → right" not in tea)
chk('Slide 9 notes key synced',"3 ✓ She doesn't work on Fridays." in notes[8] and '1 ✗' in notes[8] and '2 ✗' in notes[8] and '4 ✗' in notes[8] and '5 ✓' in notes[8])
chk('Plan key: exact 1W 2W 3R 4W 5R',all(k in tea for k in ["1 He work in a clinic. → wrong — He works in a clinic.","2 I works from Monday to Friday. → wrong — I work from Monday to Friday.","3 She doesn't work on Fridays. → right","4 Does he works on Saturdays? → wrong — Does he work on Saturdays?","5 Do you work at a hospital? → right"]))
chk('Stage 5 notes in plan mention sentence 3 right',"3 right — She doesn't work on Fridays." in tea)
chk('Other Stage 5 sentences unchanged',all(x in slides[8] for x in ['He work in a clinic.','I works from Monday to Friday.','Does he works on Saturdays?','Do you work at a hospital?']))
chk('No REV1.1 labels inside FINAL files','REV1.1' not in ' '.join(slides+notes+stu)+tea)

# language regexes (report hits)
V='work|see|finish|start|have|live|watch|examine|check|arrive|get|go|take'
pats={'3rd person w/o -s':rf"\b(?<!does )(?<!doesn't )(?<!Does )(He|he|She|she) ({V})\b",'aux + -s':rf"\b(does|Does|doesn't|Doesn't) (?:(?:he|she|it) )?({V})(e?s)\b",'I/you + -s':r"\b(I|You|you|We|They) (works|sees|finishes|starts|lives|watches)\b",'am + verb':r"\bI am (work|see|start)\b","he/she don't":r"\b(he|she|He|She) don't\b"}
hits=[]
for nm,p in pats.items():
    for k,t in enumerate(slides+stu+[tea]):
        for m in re.finditer(p,t):
            ctx=t[max(0,m.start()-25):m.end()+25].replace('\n',' ')
            hits.append((nm,('slide %d'%(k+1)) if k<17 else ('pdf p%d'%(k-16)) if k<23 else 'plan',ctx))
for r in R: print('PASS' if r[1] else 'FAIL','|',r[0],'|' if r[2] else '',r[2] if not r[1] else '')
print('--- language pattern hits (to verify as deliberate) ---')
for h in hits: print(h)
